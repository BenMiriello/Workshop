"""Converting uploaded binaries into text Claude can read, and doing it safely.

The archive cases are the point: a traversal path, a symlink and a binary entry
must never reach the model as decoded bytes, and a converted file must store
BOTH the original and its derivation and GC both.
"""

import io
import tarfile
import zipfile

import pytest

from services import attachment_convert as ac
from services import chat_attachments as ca

NS = "http://schemas.openxmlformats.org"  # OOXML namespace URIs, not network


@pytest.fixture
def store(tmp_path, monkeypatch):
    """Redirect the blob store to tmp_path so a test never writes real files."""
    fdir = tmp_path / "files"
    fdir.mkdir()
    monkeypatch.setattr(ca, "_files_dir", lambda owner_id=None: fdir)
    return fdir


def _xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["name", "qty"])
    ws.append(["apple", 3])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx():
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Hello Deck"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _docx(text):
    ct = (f'<?xml version="1.0"?><Types xmlns="{NS}/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = (f'<?xml version="1.0"?><Relationships xmlns="{NS}/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{NS}/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    doc = (f'<?xml version="1.0"?><w:document xmlns:w="{NS}/wordprocessingml/2006/main">'
           f'<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>')
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
    return b.getvalue()


def test_convertible_routing():
    assert ac.convertible("", "a.docx") == "docx"
    assert ac.convertible("application/zip", "x") == "zip"
    assert ac.convertible("", "a.tgz") == "tar"
    # A native type never routes here; kind_for already claimed it.
    assert ac.convertible("image/png", "a.png") is None
    assert ac.convertible("", "movie.mp4") is None


def test_office_conversions_produce_text():
    assert "apple" in ac.convert("xlsx", _xlsx(), "b.xlsx").text
    assert "Hello Deck" in ac.convert("pptx", _pptx(), "d.pptx").text
    assert "Report body" in ac.convert("docx", _docx("Report body text"), "r.docx").text


def test_corrupt_office_degrades_to_a_note_not_a_crash():
    r = ac.convert("docx", b"not a real docx", "broken.docx")
    assert r.text is None
    assert r.note  # says why, so the model is not handed garbage


def test_zip_skips_traversal_and_does_not_inline_binaries():
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("src/main.py", "print('hi')\n")
        z.writestr("logo.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 40)
        z.writestr("../evil.txt", "escape")
    r = ac.convert("zip", b.getvalue(), "bundle.zip")
    assert "print('hi')" in r.text          # text entry inlined
    assert "--- logo.png ---" not in r.text  # binary listed, not inlined
    assert "escape" not in r.text            # traversal entry never read
    assert any("unsafe path" in w for w in r.warnings)


def test_tar_skips_symlink_members():
    b = io.BytesIO()
    with tarfile.open(fileobj=b, mode="w:gz") as t:
        data = b"real notes\n"
        ti = tarfile.TarInfo("notes.txt")
        ti.size = len(data)
        t.addfile(ti, io.BytesIO(data))
        sl = tarfile.TarInfo("link")
        sl.type = tarfile.SYMTYPE
        sl.linkname = "/etc/passwd"
        t.addfile(sl)
    r = ac.convert("tar", b.getvalue(), "arc.tgz")
    assert "real notes" in r.text
    assert any("link" in w for w in r.warnings)


def test_zip_bomb_ratio_entry_is_not_inlined():
    # 2 MB of zeros compresses to almost nothing: a > BOMB_RATIO entry. It is
    # listed but never decoded into the derivation.
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("bomb.txt", b"\x00" * (2 * 1024 * 1024))
    r = ac.convert("zip", b.getvalue(), "z.zip")
    assert "- bomb.txt" in r.text                 # listed in the manifest
    assert "--- bomb.txt ---" not in r.text        # not inlined
    assert any("compression ratio" in w for w in r.warnings)


def test_ingest_stores_both_blobs_and_content_block_reads_the_derivation(store):
    rec = ca.ingest(_xlsx(), "book.xlsx", "")
    assert rec["kind"] == "text"
    assert rec["converter"] == "xlsx"
    assert rec["derived_sha"] and rec["sha256"] != rec["derived_sha"]
    assert rec["media_type"].endswith("spreadsheetml.sheet")
    assert ca.blob_for(rec) is not None       # original preserved
    assert ca.derived_path(rec) is not None   # derivation stored

    blocks = ca.content_blocks([{"id": rec["id"], "name": rec["name"]}])
    assert len(blocks) == 1
    assert blocks[0]["type"] == "text"
    assert "apple" in blocks[0]["text"]
    assert "converted to text" in blocks[0]["text"]


def test_delete_gcs_both_blobs(store):
    rec = ca.ingest(_xlsx(), "book.xlsx", "")
    osha, dsha = rec["sha256"], rec["derived_sha"]
    assert ca.delete(rec["id"]) is True
    assert not ca._blob_path(None, osha).exists()
    assert not ca._blob_path(None, dsha).exists()


def test_native_text_is_unchanged_by_the_new_path(store):
    rec = ca.ingest(b"hello world\n", "notes.txt", "text/plain")
    assert rec["kind"] == "text"
    assert "converter" not in rec
    block = ca.content_blocks([{"id": rec["id"], "name": rec["name"]}])[0]
    assert block["text"].startswith("Attached file: notes.txt")


def test_unsupported_type_still_refused(store):
    with pytest.raises(ValueError, match="not a type a chat can read"):
        ca.ingest(b"\x00\x01\x02", "mystery.bin", "application/octet-stream")


def test_encrypted_or_corrupt_convertible_stores_original_with_a_note(store):
    # A ".docx" that is not a valid zip: stored (original kept), no derivation,
    # and the model block says it could not be read rather than inventing text.
    rec = ca.ingest(b"this is not a docx", "broken.docx", "")
    assert rec["converter"] == "docx"
    assert "derived_sha" not in rec
    assert ca.blob_for(rec) is not None
    block = ca.content_blocks([{"id": rec["id"], "name": rec["name"]}])[0]
    assert "could not be read" in block["text"]
    assert "original is stored" in block["text"]
